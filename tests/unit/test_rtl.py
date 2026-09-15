from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from docslides.slides.rtl import apply_rtl_to_slide, set_textframe_rtl

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _make_slide_with_title():
    prs = Presentation()
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text_frame.text = "שלום עולם"
    return prs, slide


def test_set_textframe_rtl_sets_ooxml_attribute_and_alignment():
    prs, slide = _make_slide_with_title()
    text_frame = slide.shapes.title.text_frame

    set_textframe_rtl(text_frame)

    for paragraph in text_frame.paragraphs:
        pPr = paragraph._p.find(f"{{{A_NS}}}pPr")
        assert pPr is not None
        assert pPr.get("rtl") == "1"
        assert paragraph.alignment == PP_ALIGN.RIGHT


def test_apply_rtl_to_slide_covers_all_text_shapes():
    prs, slide = _make_slide_with_title()

    apply_rtl_to_slide(slide, prs.slide_width)

    title_pPr = slide.shapes.title.text_frame.paragraphs[0]._p.find(f"{{{A_NS}}}pPr")
    assert title_pPr is not None
    assert title_pPr.get("rtl") == "1"


def test_mixed_direction_deck_only_affects_targeted_slides():
    prs = Presentation()
    layout = prs.slide_layouts[0]

    english_slide = prs.slides.add_slide(layout)
    english_slide.shapes.title.text_frame.text = "Hello world"

    hebrew_slide = prs.slides.add_slide(layout)
    hebrew_slide.shapes.title.text_frame.text = "שלום עולם"
    apply_rtl_to_slide(hebrew_slide, prs.slide_width)

    english_pPr = english_slide.shapes.title.text_frame.paragraphs[0]._p.find(f"{{{A_NS}}}pPr")
    hebrew_pPr = hebrew_slide.shapes.title.text_frame.paragraphs[0]._p.find(f"{{{A_NS}}}pPr")

    assert english_pPr is None or english_pPr.get("rtl") is None
    assert hebrew_pPr is not None and hebrew_pPr.get("rtl") == "1"
