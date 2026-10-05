"""Generated files: the builders render Gemma 4's specs into files that open, and the chat routes
a request for a file to the generator."""

from __future__ import annotations

import asyncio

import pymupdf
import pytest
from openpyxl import load_workbook
from pptx import Presentation

from docslides.api import routes_chat as rc
from docslides.documents.generator import build_document
from docslides.documents.pdf_builder import build_pdf
from docslides.documents.xlsx_builder import build_xlsx
from docslides.llm.schemas import (
    ChatIntent,
    PdfSection,
    PdfSpec,
    PresentationSpec,
    SheetSpec,
    SlideContent,
    SpreadsheetSpec,
    TableSpec,
)

BUDGET = SpreadsheetSpec(
    title="Budget",
    sheets=[
        SheetSpec(
            name="Q1: costs/plan",
            columns=["Item", "Cost (USD)"],
            rows=[["Rent", 1200.0], ["Food", "1,250"], ["Total", "=SUM(B2:B3)"]],
        ),
        SheetSpec(name="Q1: costs/plan", columns=["Note"], rows=[["second sheet, same name"]]),
    ],
)


def test_a_workbook_keeps_numbers_formulas_and_valid_unique_sheet_names(tmp_path):
    path = build_xlsx(BUDGET, tmp_path / "budget.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Q1  costs plan", "Q1  costs plan (2)"]
    sheet = workbook.worksheets[0]
    assert [c.value for c in sheet[1]] == ["Item", "Cost (USD)"]
    assert sheet["B2"].value == 1200 and sheet["B3"].value == 1250
    assert sheet["B4"].value == "=SUM(B2:B3)"
    assert sheet.freeze_panes == "A2" and sheet["A1"].font.bold


def test_a_hebrew_workbook_reads_right_to_left(tmp_path):
    path = build_xlsx(BUDGET, tmp_path / "budget.xlsx", lang="he")
    assert load_workbook(path).worksheets[0].sheet_view.rightToLeft


REPORT = PdfSpec(
    title="Quarterly <Report>",
    subtitle="Q1 2026",
    sections=[
        PdfSection(
            heading="Summary",
            paragraphs=["Revenue grew & costs fell."],
            bullets=["Point one", "Point two"],
            table=TableSpec(columns=["Metric", "Value"], rows=[["Revenue", "10"], ["Short row"]]),
        )
    ],
)


def test_a_pdf_has_the_title_sections_lists_and_table(tmp_path):
    path = build_pdf(REPORT, tmp_path / "report.pdf")
    with pymupdf.open(str(path)) as doc:
        text = "".join(page.get_text() for page in doc)
        assert doc.metadata["title"] == "Quarterly <Report>"
    for expected in ("Quarterly <Report>", "Revenue grew & costs fell.", "Point two", "Metric", "Short row"):
        assert expected in text


def test_a_long_hebrew_pdf_flows_onto_more_pages(tmp_path):
    spec = PdfSpec(
        title="דוח",
        sections=[PdfSection(heading=f"סעיף {i}", paragraphs=["שלום עולם " * 80]) for i in range(12)],
    )
    path = build_pdf(spec, tmp_path / "he.pdf", lang="he")
    with pymupdf.open(str(path)) as doc:
        assert doc.page_count > 1
        assert "שלום" in doc[0].get_text()


def test_a_presentation_gets_a_title_slide_before_its_slides(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "docslides.documents.generator.get_config",
        lambda: type("C", (), {"paths": type("P", (), {"output_dir": str(tmp_path)})()})(),
    )
    spec = PresentationSpec(
        title="Solar Energy",
        subtitle="An overview",
        slides=[SlideContent(title="Why solar", bullets=["Cheap", "Clean"], layout_type="title_bullets")],
    )
    path = build_document(spec, "pptx", "en", "job1")
    assert path.name == "Solar_Energy_job1.pptx"
    titles = [slide.shapes.title.text for slide in Presentation(str(path)).slides]
    assert titles == ["Solar Energy", "Why solar"]


@pytest.mark.parametrize(
    ("intent", "has_attachment", "expected"),
    [
        (ChatIntent(wants_slides=False, document_format="xlsx"), False, "xlsx"),
        (ChatIntent(wants_slides=False, document_format="pdf"), True, "pdf"),
        (ChatIntent(wants_slides=True), False, "pptx"),
        # A deck from an attached document goes through the full slide pipeline instead.
        (ChatIntent(wants_slides=True, document_format="pptx"), True, None),
        (ChatIntent(wants_slides=False), False, None),
    ],
)
def test_which_requests_generate_a_file(intent, has_attachment, expected):
    assert rc._requested_file(intent, has_attachment) == expected


@pytest.mark.parametrize(
    ("message", "expected"),
    [
        (
            "create an excel sheet to calculate the daily income and calculate the VAT, in addition to "
            "tips which do not add to the VAT and calculated at the end of the month",
            "xlsx",
        ),
        ("Make me a spreadsheet for my monthly budget", "xlsx"),
        ("please generate a PDF report on solar energy", "pdf"),
        ("How do I sum a column in Excel?", None),
        ("Summarize this:\n" + "x " * 200 + "create an excel sheet", None),
    ],
)
def test_explicit_file_request_fallback(message, expected):
    assert rc._explicit_file_format(message) == expected


class FakeGemma:
    def __init__(self, intent: ChatIntent) -> None:
        self.intent = intent
        self.spec_prompts: list[list] = []

    async def complete_json(self, messages, call_site, schema, **_):
        if schema is ChatIntent:
            return self.intent
        assert schema is SpreadsheetSpec and call_site.name == "document_generation"
        self.spec_prompts.append(messages)
        return BUDGET


def test_a_chat_request_for_a_spreadsheet_returns_a_download(tmp_path, monkeypatch):
    events: list[tuple[str, dict]] = []

    async def publish_status(job_id, message, **_):
        events.append(("status", {"message": message}))

    async def publish_done(job_id, **extra):
        events.append(("done", extra))

    async def publish_error(job_id, message):
        events.append(("error", {"message": message}))

    monkeypatch.setattr(rc.event_bus, "publish_status", publish_status)
    monkeypatch.setattr(rc.event_bus, "publish_done", publish_done)
    monkeypatch.setattr(rc.event_bus, "publish_error", publish_error)
    monkeypatch.setattr(rc, "detect_language", lambda text: "en")
    monkeypatch.setattr(
        rc, "build_document", lambda spec, fmt, lang, job_id: build_xlsx(spec, tmp_path / f"{job_id}.xlsx")
    )
    gemma = FakeGemma(ChatIntent(wants_slides=False, document_format="xlsx", target_lang="fr"))
    monkeypatch.setattr(rc, "get_client", lambda: gemma)

    req = rc.ChatRequest(messages=[{"role": "user", "content": "Make me an Excel budget in French"}])
    asyncio.run(rc._run_chat_turn("job-x", req))

    assert events[-1] == ("done", {"output_path": str(tmp_path / "job-x.xlsx")})
    assert rc.job_outputs["job-x"] == str(tmp_path / "job-x.xlsx")
    system = gemma.spec_prompts[0][0].content
    assert "Excel workbook" in system and "in French" in system and "Clara" in system


def test_an_rtl_table_puts_its_first_column_on_the_right():
    from docslides.documents.pdf_builder import _html

    spec = PdfSpec(
        title="t",
        sections=[PdfSection(heading="h", table=TableSpec(columns=["A", "B"], rows=[["1", "2"]]))],
    )
    assert "<th>B</th><th>A</th>" in _html(spec, rtl=True)
    assert "<th>A</th><th>B</th>" in _html(spec, rtl=False)
