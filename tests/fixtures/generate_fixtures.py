#!/usr/bin/env python3
"""Generate the three test fixtures used to validate OCR-detection and
language-routing logic:

  1. scanned_arabic.pdf   -- an image-only PDF page (no text layer) with
                              Arabic text rendered into it, so
                              ingestion/parser.py's scan-detection heuristic
                              must flag it SCANNED and route it to the
                              Arabic OCR engine.
  2. native_hebrew.docx   -- a DOCX with real Hebrew Unicode text in the
                              document body (native text layer, no OCR
                              needed), to validate Hebrew language detection
                              and RTL handling downstream.
  3. mixed_language.docx  -- a DOCX with English, French, and Arabic
                              paragraphs, to validate that language
                              detection and chunking operate per
                              segment/page rather than assuming one language
                              for the whole document.

Run once: `python tests/fixtures/generate_fixtures.py`. Requires reportlab,
Pillow, and python-docx (see pyproject.toml's `dev` extra).

Note: Arabic glyphs are rendered without contextual shaping (no `libraqm` in
the Pillow build assumed here), so letters appear in isolated form rather
than properly joined. That's fine for these fixtures -- they exist to
exercise the scan-detection/language-routing pipeline, not to assert OCR
transcription accuracy.
"""

from __future__ import annotations

import io
from pathlib import Path

FIXTURES_DIR = Path(__file__).parent

WINDOWS_FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/arial.ttf"),
    Path("C:/Windows/Fonts/tahoma.ttf"),
    Path("C:/Windows/Fonts/david.ttf"),
]
LINUX_FONT_CANDIDATES = [
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf"),
]


def _find_font() -> Path:
    for candidate in WINDOWS_FONT_CANDIDATES + LINUX_FONT_CANDIDATES:
        if candidate.exists():
            return candidate
    raise RuntimeError(
        "No Unicode-capable TTF font found. Install one (e.g. DejaVu Sans or "
        "Noto Sans Arabic) and add its path to generate_fixtures.py."
    )


def generate_scanned_arabic_pdf() -> Path:
    from PIL import Image, ImageDraw, ImageFont
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    font_path = _find_font()
    font = ImageFont.truetype(str(font_path), 28)

    image = Image.new("RGB", (1600, 2200), "white")
    draw = ImageDraw.Draw(image)
    arabic_lines = [
        "هذا مستند تجريبي باللغة العربية",
        "يستخدم لاختبار كشف الصفحات الممسوحة ضوئياً",
        "والتوجيه الصحيح لمحرك التعرف الضوئي على الحروف",
    ]
    y = 150
    for line in arabic_lines:
        draw.text((100, y), line, font=font, fill="black")
        y += 80

    img_buffer = io.BytesIO()
    image.save(img_buffer, format="PNG")
    img_buffer.seek(0)

    output_path = FIXTURES_DIR / "scanned_arabic.pdf"
    c = canvas.Canvas(str(output_path), pagesize=LETTER)
    # Draw the rasterized text as an image -- this PDF page has NO text
    # layer, which is exactly what the scan-detection heuristic must catch.
    from reportlab.lib.utils import ImageReader

    c.drawImage(ImageReader(image), 0, 0, width=LETTER[0], height=LETTER[1])
    c.showPage()
    c.save()
    return output_path


def generate_native_hebrew_docx() -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("מסמך בדיקה", level=1)
    doc.add_paragraph(
        "זהו מסמך טקסט טבעי בעברית המשמש לבדיקת זיהוי שפה וניתוב טקסט מימין "
        "לשמאל, ללא צורך בזיהוי אופטי של תווים."
    )
    doc.add_paragraph(
        "פסקה שנייה לבדיקת חלוקה למקטעים (chunking) ושמירה על גבולות המשפט."
    )

    output_path = FIXTURES_DIR / "native_hebrew.docx"
    doc.save(str(output_path))
    return output_path


def generate_mixed_language_docx() -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Mixed-language test document", level=1)
    doc.add_paragraph(
        "This is a paragraph written in English, used to validate that "
        "language detection operates per segment rather than assuming a "
        "single language for the whole document."
    )
    doc.add_paragraph(
        "Ceci est un paragraphe rédigé en français, afin de vérifier que le "
        "changement de langue au sein d'un même document est correctement "
        "détecté."
    )
    doc.add_paragraph(
        "هذه فقرة مكتوبة باللغة العربية للتحقق من التوجيه الصحيح للنص من "
        "اليمين إلى اليسار ضمن مستند متعدد اللغات."
    )

    output_path = FIXTURES_DIR / "mixed_language.docx"
    doc.save(str(output_path))
    return output_path


def main() -> None:
    paths = [
        generate_scanned_arabic_pdf(),
        generate_native_hebrew_docx(),
        generate_mixed_language_docx(),
    ]
    for path in paths:
        print(f"wrote {path}")


if __name__ == "__main__":
    main()
