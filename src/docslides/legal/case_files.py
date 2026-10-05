"""The attorney's own case files: chunked, embedded and searched so the user can chat with them.
Driven by scripts/ingest_case_files.py; settings under legal.case_files in config/config.yaml.

Files go in legal_data/ at the project root, one subfolder per matter (client / case), any depth
below it:

    legal_data/
      Cohen v. Levi/
        pleadings/statement_of_claim.pdf
        correspondence/2024-03-01 letter to opposing counsel.docx
        hearing 2024-05-12 transcript.pdf
      Estate of Mizrahi/
        ...

The top-level subfolder is the chunk's `matter`, so a chat can be scoped to one case. Files
directly in legal_data/ get the matter "(general)".

Chunking strategy, and why (case files are not statutes, and not all case files are alike):

  1. Structure first, size second. A lawyer cites "paragraph 14 of the statement of claim",
     "clause 7.2 of the agreement", "p. 23 of the transcript". So the unit of meaning is the
     numbered paragraph / clause / speaker turn / e-mail, never an arbitrary token window: units
     are never split unless one alone exceeds the budget (then by sentence), and chunks are packed
     from whole units up to ~450 tokens -- small enough for precise retrieval, big enough to keep
     an allegation together with its supporting facts.
  2. Document-type aware. Each file is classified (pleading, judgment/decision, affidavit,
     contract, transcript, e-mail, correspondence, evidence/exhibit, memo/notes, other) from its
     name and opening text, Hebrew and English, and the type picks the splitter:
       * pleadings, judgments, affidavits, memos: numbered paragraphs ("12.", "(א)", "סעיף 5")
         and headings; a heading starts a new chunk and stays with the chunk as its section;
       * contracts: clauses and sub-clauses ("7", "7.2", "(b)") kept whole;
       * transcripts / protocols: speaker turns, and a question is kept with its answer
         ("ש:" + "ת:", "Q." + "A.") so testimony is never cut between them;
       * e-mail threads: one message at a time -- a chunk never spans two messages -- and each
         chunk repeats its message's sender, recipients, date and subject;
       * everything else (letters, exhibits, invoices, spreadsheets): paragraphs packed by size.
  3. Every chunk carries a context header -- matter, file path, document type, document date,
     page range, section heading, paragraph numbers -- and the header is embedded with the text.
     Case files are full of near-identical passages (the claim and the defence describe the same
     events); the header is what lets "what does the defence say about the delivery date" find the
     defence's paragraph and not the claim's, and it is what the answer cites.
  4. Page numbers are kept per chunk (PDF pages; OCRed scans too), so answers cite "p. 4".
  5. Small overlap (~60 tokens, whole units only) between consecutive chunks of the same
     document, so a fact split across a chunk boundary is still retrievable.
  6. Hybrid retrieval: dense (bge-m3, multilingual) + BM25 (names, case numbers, amounts, dates
     and terms of art, which embeddings rank loosely), fused, reranked by the cross-encoder, then
     the paragraphs right before and after the best hits are added ("small-to-big"), so the model
     sees an allegation's context without the index holding oversized chunks.

Scanned pages and images are OCRed (ocr/router.py) when an OCR engine is installed; otherwise
they are skipped and named in the run's report, never silently dropped.
"""

from __future__ import annotations

import asyncio
import email
import email.policy
import hashlib
import html
import re
from dataclasses import dataclass, field
from datetime import date
from functools import lru_cache
from pathlib import Path

from docslides.cleaning.tokens import count_tokens
from docslides.config import get_config
from docslides.legal_data.corpus_chunking import CorpusChunk
from docslides.legal_data.hebrew import normalize_for_embedding, normalize_for_index
from docslides.logging_setup import get_logger

logger = get_logger(__name__)

# Bump when the chunking changes: every file is then re-chunked on the next run.
CHUNKER_VERSION = "1"
CATEGORY = "case_files"
GENERAL_MATTER = "(general)"
STATE_FILE = "_case_files_state.sqlite"

TEXT_SUFFIXES = {".txt", ".md"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
SUPPORTED_SUFFIXES = {".pdf", ".docx", ".eml", ".xlsx", *TEXT_SUFFIXES, *IMAGE_SUFFIXES}
# Formats a lawyer's folder often holds that need converting first.
CONVERT_FIRST = {".doc": ".docx", ".rtf": ".docx", ".msg": ".eml (or PDF)", ".pages": ".docx", ".odt": ".docx"}

DOC_TYPES = ("email", "transcript", "judgment", "affidavit", "pleading", "contract", "correspondence",
             "evidence", "notes", "document")

# --- classification ---------------------------------------------------------------------------

# Checked in this order against the file name and the document's opening text; the first match wins.
_TYPE_KEYWORDS: list[tuple[str, tuple[str, ...]]] = [
    ("transcript", ("פרוטוקול", "תמליל", "transcript", "deposition", "protocol", "hearing minutes")),
    ("judgment", ("פסק דין", "פסק-דין", "פסה\"ד", "גזר דין", "גזר-דין", "החלטה", "judgment", "judgement",
                  "decision", "ruling", "verdict", "court order")),
    ("affidavit", ("תצהיר", "affidavit", "declaration", "sworn statement")),
    ("pleading", ("כתב תביעה", "כתב הגנה", "כתב תשובה", "כתב טענות", "כתב ערעור", "הודעת ערעור", "בקשה",
                  "בקשת", "תגובה", "סיכומים", "סיכומי", "עתירה", "כתב אישום", "הודעה לבית", "complaint",
                  "statement of claim", "statement of defense", "statement of defence", "motion", "brief",
                  "petition", "appeal", "reply", "response to", "summation", "pleading", "indictment")),
    ("contract", ("הסכם", "חוזה", "הסכמי", "נספח להסכם", "agreement", "contract", "lease", "addendum",
                  "memorandum of understanding", "mou", "terms and conditions")),
    ("correspondence", ("מכתב", "לכבוד", "הנדון", "התראה", "letter", "dear ", "re:", "notice", "demand")),
    ("evidence", ("נספח", "מוצג", "חשבונית", "קבלה", "דוח", "דו\"ח", "חוות דעת", "exhibit", "invoice",
                  "receipt", "report", "expert opinion", "statement of account", "appendix", "annex")),
    ("notes", ("תזכיר", "סיכום פגישה", "הערות", "רשימות", "memo", "memorandum", "notes", "meeting summary",
               "file note", "attendance note")),
]
_EMAIL_HEADER_RE = re.compile(r"^\s*(From|To|Cc|Sent|Date|Subject|מאת|אל|עותק|נשלח|תאריך|נושא)\s*:",
                              re.IGNORECASE | re.MULTILINE)
_QA_LINE_RE = re.compile(r"^\s*(?:ש|ת|Q|A|שאלה|תשובה)\s*[.:]\s", re.MULTILINE)


def classify(name: str, text: str) -> str:
    """The document type, from its file name and opening text (Hebrew or English)."""
    head = text[:400].lower()  # the title area: further down, any word can appear
    lowered_name = name.lower()
    if lowered_name.endswith(".eml") or len(_EMAIL_HEADER_RE.findall(text[:800])) >= 3:
        return "email"
    if len(_QA_LINE_RE.findall(text[:6000])) >= 6:
        return "transcript"
    for doc_type, keywords in _TYPE_KEYWORDS:  # the file name says most, so it is checked first
        if any(k in lowered_name for k in keywords):
            return doc_type
    for doc_type, keywords in _TYPE_KEYWORDS:
        if any(k in head for k in keywords):
            return doc_type
    return "document"


# --- dates ------------------------------------------------------------------------------------

_HEBREW_MONTHS = {
    "ינואר": 1, "פברואר": 2, "מרץ": 3, "מרס": 3, "אפריל": 4, "מאי": 5, "יוני": 6, "יולי": 7,
    "אוגוסט": 8, "ספטמבר": 9, "אוקטובר": 10, "נובמבר": 11, "דצמבר": 12,
}
_ENGLISH_MONTHS = {m: i for i, m in enumerate(
    ("january", "february", "march", "april", "may", "june", "july", "august", "september", "october",
     "november", "december"), 1)}
_ENGLISH_MONTHS.update({m[:3]: i for m, i in list(_ENGLISH_MONTHS.items())})
_MONTH_NAMES = "|".join(sorted([*_HEBREW_MONTHS, *_ENGLISH_MONTHS], key=len, reverse=True))
_DATE_PATTERNS = [
    # 2024-03-01
    (re.compile(r"\b((?:19|20)\d\d)-(\d{1,2})-(\d{1,2})\b"), lambda m: (m[1], m[2], m[3])),
    # 01/03/2024, 1.3.2024, 01-03-2024, 01/03/24 -- Israeli/European order: day first. A two-digit
    # year only with slashes: "1.2.10" is far more often a clause number than a date.
    (re.compile(r"\b(\d{1,2})[./-](\d{1,2})[./-]((?:19|20)\d\d)\b"), lambda m: (m[3], m[2], m[1])),
    (re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d\d)\b"), lambda m: (m[3], m[2], m[1])),
    # 1 במרץ 2024, 1 March 2024
    (re.compile(rf"\b(\d{{1,2}})\s+ב?({_MONTH_NAMES})\.?,?\s+((?:19|20)\d\d)\b", re.IGNORECASE),
     lambda m: (m[3], m[2], m[1])),
    # March 1, 2024
    (re.compile(rf"\b({_MONTH_NAMES})\.?\s+(\d{{1,2}}),?\s+((?:19|20)\d\d)\b", re.IGNORECASE),
     lambda m: (m[3], m[1], m[2])),
]


def _to_date(year: str, month: str, day: str) -> date | None:
    try:
        y = int(year) if len(year) == 4 else 2000 + int(year)
        mo = (int(month) if month.isdigit()
              else _HEBREW_MONTHS.get(month) or _ENGLISH_MONTHS.get(month.lower().rstrip(".")))
        return date(y, mo, int(day)) if mo else None
    except (ValueError, TypeError):
        return None


def find_date(text: str) -> str:
    """The first plausible date in `text`, as YYYY-MM-DD, or "" if there is none."""
    found: list[tuple[int, date]] = []
    for pattern, parts in _DATE_PATTERNS:
        for match in pattern.finditer(text):
            parsed = _to_date(*parts(match))
            if parsed and 1950 <= parsed.year <= 2100:
                found.append((match.start(), parsed))
                break
    return min(found)[1].isoformat() if found else ""


# --- extraction -------------------------------------------------------------------------------


@dataclass
class PageText:
    number: int | None  # 1-based page; None for formats without pages (DOCX, TXT, e-mail)
    text: str
    wrapped: bool = True  # lines are visual lines (PDF/OCR/hard-wrapped text), not paragraphs


@dataclass
class CaseDocument:
    path: Path
    rel_path: str
    matter: str
    doc_type: str
    title: str
    doc_date: str
    pages: list[PageText]
    notes: list[str] = field(default_factory=list)  # pages skipped, attachments not read, ...


def matter_of(rel_path: str) -> str:
    parts = Path(rel_path).parts
    return parts[0] if len(parts) > 1 else GENERAL_MATTER


def is_case_file(path: Path) -> bool:
    name = path.name
    return (path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES and not name.startswith((".", "~$"))
            and not name.endswith(".meta.json"))


def list_case_files(folder: Path) -> tuple[list[Path], list[Path]]:
    """(files to index, files skipped because their format must be converted first)."""
    supported, unsupported = [], []
    for path in sorted(folder.rglob("*")):
        relative = path.relative_to(folder)
        if any(part.startswith(".") for part in relative.parts) or relative.as_posix() == "README.md":
            continue  # hidden files, and the folder's own instructions
        if is_case_file(path):
            supported.append(path)
        elif path.is_file() and path.suffix.lower() in CONVERT_FIRST:
            unsupported.append(path)
    return supported, unsupported


def _ocr_png(png: bytes, lang: str, page_index: int) -> str:
    from docslides.ocr.router import recognize_page

    return asyncio.run(recognize_page(png, lang, page_index=page_index)).text


def _ocr_pages(path: Path, numbers: list[int], lang: str, notes: list[str]) -> dict[int, str]:
    """OCR text of the given 1-based pages (of a PDF or image), or {} with a note when OCR is off or
    unavailable."""
    if not numbers:
        return {}
    shown = ", ".join(map(str, numbers[:20])) + ("..." if len(numbers) > 20 else "")
    if not get_config().legal.case_files.ocr_scanned_pages:
        notes.append(f"scanned page(s) {shown} skipped (ocr_scanned_pages is off)")
        return {}
    import pymupdf

    out: dict[int, str] = {}
    with pymupdf.open(path) as doc:
        for number in numbers:
            try:
                png = doc[number - 1].get_pixmap(dpi=200).tobytes("png")
                out[number] = _ocr_png(png, lang, number - 1)
            except Exception as exc:  # noqa: BLE001 -- no engine installed, or this page failed
                notes.append(f"scanned page(s) {shown} not read: OCR failed ({exc})")
                logger.warning("case_files_ocr_failed", path=str(path), page=number, error=str(exc))
                return out
    return out


def _document_language(text: str) -> str:
    try:
        from docslides.ingestion.language_detect import detect_language

        return detect_language(text[:3000]) or get_config().legal.case_files.ocr_default_language
    except Exception:  # noqa: BLE001 -- the detector's models are optional
        return get_config().legal.case_files.ocr_default_language


def _pdf_pages(path: Path, notes: list[str]) -> list[PageText]:
    from docslides.legal.pdf_text import MARGIN_TITLE_CLOSE, MARGIN_TITLE_OPEN, extract_pdf_pages

    try:
        texts, scanned = extract_pdf_pages(path)
    except Exception as exc:  # noqa: BLE001 -- the layout rebuild is tuned for gazettes; plain text works
        logger.warning("case_files_pdf_rebuild_failed", path=str(path), error=str(exc))
        import pymupdf

        with pymupdf.open(path) as doc:
            texts = [page.get_text("text", sort=True) for page in doc]
        scanned = [i for i, t in enumerate(texts, 1) if len(re.findall(r"\w", t)) < 20]
    texts = [t.replace(MARGIN_TITLE_OPEN, "").replace(MARGIN_TITLE_CLOSE, "") for t in texts]
    from docslides.legal.sources import looks_visual_order

    if looks_visual_order("\n".join(texts)):
        from docslides.legal_data.hebrew import visual_to_logical

        texts = ["\n".join(visual_to_logical(line) for line in t.splitlines()) for t in texts]
        notes.append("Hebrew text layer was in visual (reversed) order; lines were reversed -- check quotes")
    ocr = _ocr_pages(path, scanned, _document_language("\n".join(texts)), notes)
    return [PageText(n, ocr.get(n, text) if n in scanned else text) for n, text in enumerate(texts, 1)]


def _image_pages(path: Path, notes: list[str]) -> list[PageText]:
    import pymupdf

    with pymupdf.open(path) as doc:
        count = doc.page_count
    ocr = _ocr_pages(path, list(range(1, count + 1)), get_config().legal.case_files.ocr_default_language, notes)
    return [PageText(n, ocr.get(n, "")) for n in range(1, count + 1)]


_LIST_FORMATS = {
    "decimal": lambda n: f"{n}.",
    "hebrew1": lambda n: f"{'אבגדהוזחטיכלמנסעפצקרשת'[(n - 1) % 22]}.",
    "hebrew2": lambda n: f"{'אבגדהוזחטיכלמנסעפצקרשת'[(n - 1) % 22]}.",
    "lowerLetter": lambda n: f"({chr(96 + (n - 1) % 26 + 1)})",
    "upperLetter": lambda n: f"({chr(64 + (n - 1) % 26 + 1)})",
    "lowerRoman": lambda n: f"({['i', 'ii', 'iii', 'iv', 'v', 'vi', 'vii', 'viii', 'ix', 'x'][(n - 1) % 10]})",
}


def _numbering_formats(document) -> dict[tuple[str, str], str]:
    """(numId, ilvl) -> numFmt from the DOCX's numbering part, so auto-numbered paragraphs (Word's
    own "1." / "(א)" lists, which never appear in the paragraph text) get their numbers back --
    a pleading's paragraph numbers are what a lawyer cites."""
    try:
        root = document.part.numbering_part.element
    except Exception:  # noqa: BLE001 -- no numbering part
        return {}
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    abstract = {}
    for node in root.iter(f"{w}abstractNum"):
        levels = {}
        for lvl in node.iter(f"{w}lvl"):
            fmt = lvl.find(f"{w}numFmt")
            levels[lvl.get(f"{w}ilvl")] = fmt.get(f"{w}val") if fmt is not None else "decimal"
        abstract[node.get(f"{w}abstractNumId")] = levels
    out = {}
    for num in root.iter(f"{w}num"):
        ref = num.find(f"{w}abstractNumId")
        if ref is not None:
            for ilvl, fmt in abstract.get(ref.get(f"{w}val"), {}).items():
                out[(num.get(f"{w}numId"), ilvl)] = fmt
    return out


def _num_pr(element, style, w: str):
    """A paragraph's list numbering: its own, else its style's (Word's "List Number" style), else
    the style's base style's."""
    num = element.find(f"{w}pPr/{w}numPr")
    while num is None and style is not None:
        num = style.element.find(f"{w}pPr/{w}numPr")
        style = style.base_style
    return num


def _docx_pages(path: Path) -> list[PageText]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    formats = _numbering_formats(document)
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    counters: dict[tuple[str, int], int] = {}
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == f"{w}p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                lines.append("")
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            if style.lower().startswith(("heading", "title", "כותרת")):
                lines += ["", f"## {text}"]
                continue
            num = _num_pr(child, paragraph.style, w)
            if num is not None and num.find(f"{w}numId") is not None:
                num_id = num.find(f"{w}numId").get(f"{w}val")
                ilvl_node = num.find(f"{w}ilvl")
                ilvl = int(ilvl_node.get(f"{w}val")) if ilvl_node is not None else 0
                fmt = formats.get((num_id, str(ilvl)), "decimal")
                if fmt in _LIST_FORMATS:
                    counters[(num_id, ilvl)] = counters.get((num_id, ilvl), 0) + 1
                    for deeper in [k for k in counters if k[0] == num_id and k[1] > ilvl]:
                        del counters[deeper]
                    text = f"{_LIST_FORMATS[fmt](counters[(num_id, ilvl)])} {text}"
                elif fmt == "bullet":
                    text = f"• {text}"
            lines.append(text)
        elif child.tag == f"{w}tbl":
            for row in Table(child, document).rows:
                cells = list(dict.fromkeys(cell.text.strip() for cell in row.cells))  # merged cells repeat
                if any(cells):
                    lines.append(" | ".join(cells))
            lines.append("")
    return [PageText(None, "\n".join(lines), wrapped=False)]


def _xlsx_pages(path: Path) -> list[PageText]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines += ["", f"## {sheet.title}"]
        for row in sheet.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
    workbook.close()
    return [PageText(None, "\n".join(lines), wrapped=False)]


def _strip_html(text: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?</\1>", "", text)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>", "\n", text)
    return html.unescape(re.sub(r"<[^>]+>", "", text))


def _eml_pages(path: Path, notes: list[str]) -> list[PageText]:
    message = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    header = [f"{name}: {message[name]}" for name in ("From", "To", "Cc", "Date", "Subject") if message[name]]
    body = message.get_body(preferencelist=("plain", "html"))
    text = body.get_content() if body is not None else ""
    if body is not None and body.get_content_type() == "text/html":
        text = _strip_html(text)
    attachments = [part.get_filename() for part in message.iter_attachments() if part.get_filename()]
    if attachments:
        notes.append("attachments not indexed (save them into the folder to index them): " + ", ".join(attachments))
    return [PageText(None, "\n".join(header) + "\n\n" + text, wrapped=not any(len(x) > 300 for x in text.splitlines()))]


def _text_pages(path: Path) -> list[PageText]:
    text = path.read_text(encoding="utf-8", errors="replace")
    # Form feeds mark pages in text exported from PDFs.
    pages = text.split("\f")
    wrapped = not any(len(line) > 300 for line in text.splitlines())
    if len(pages) > 1:
        return [PageText(n, t, wrapped) for n, t in enumerate(pages, 1)]
    return [PageText(None, text, wrapped)]


def _title(pages: list[PageText], fallback: str) -> str:
    for page in pages[:2]:
        for line in page.text.splitlines():
            line = re.sub(r"^#+\s*|\s+", " ", line).strip()
            if 4 <= len(line) <= 120 and re.search(r"\w{3}", line) and not _EMAIL_HEADER_RE.match(line):
                return line
    return fallback


def load_document(path: Path, root: Path) -> CaseDocument:
    """A case file's text, page by page, with its matter, type, title and date."""
    suffix = path.suffix.lower()
    notes: list[str] = []
    if suffix == ".pdf":
        pages = _pdf_pages(path, notes)
    elif suffix == ".docx":
        pages = _docx_pages(path)
    elif suffix == ".xlsx":
        pages = _xlsx_pages(path)
    elif suffix == ".eml":
        pages = _eml_pages(path, notes)
    elif suffix in IMAGE_SUFFIXES:
        pages = _image_pages(path, notes)
    elif suffix in TEXT_SUFFIXES:
        pages = _text_pages(path)
    else:
        raise ValueError(f"unsupported file type {suffix}")
    rel_path = path.relative_to(root).as_posix()
    opening = "\n".join(p.text for p in pages)[:6000]
    doc_type = classify(path.name, opening)
    if doc_type == "email":
        dated = re.search(r"^\s*(?:Date|Sent|תאריך|נשלח)\s*:(.*)$", opening, re.IGNORECASE | re.MULTILINE)
        doc_date = find_date(dated.group(1)) if dated else ""
        doc_date = doc_date or _email_header_date(opening)
    else:
        doc_date = find_date(path.stem) or find_date(opening[:2500])
    return CaseDocument(path=path, rel_path=rel_path, matter=matter_of(rel_path), doc_type=doc_type,
                        title=_title(pages, path.stem), doc_date=doc_date, pages=pages, notes=notes)


def _email_header_date(text: str) -> str:
    """RFC 2822 dates ("Tue, 5 Mar 2024 10:00:00 +0200") as YYYY-MM-DD."""
    from email.utils import parsedate_to_datetime

    match = re.search(r"^\s*(?:Date|Sent)\s*:\s*(.+)$", text, re.IGNORECASE | re.MULTILINE)
    if not match:
        return ""
    try:
        return parsedate_to_datetime(match.group(1).strip()).date().isoformat()
    except (TypeError, ValueError):
        return ""


# --- units: paragraphs, clauses, speaker turns ------------------------------------------------

# A line that starts a new unit: "12.", "12)", "(12)", "7.2", "(א)", "א.", "(b)", "b.", bullets,
# "סעיף 5", "Section 5". Not a bare "12": a wrapped line can start with a number ("1 March 2024").
_NUMBERED_RE = re.compile(
    r"^\s*(?P<label>\(\d{1,3}\)|\d{1,3}(?:\.\d{1,3}){0,3}[.)]|\d{1,3}(?:\.\d{1,3}){1,3}|\([א-ת]{1,2}\)|[א-ת][.)]|\([a-zA-Z]{1,4}\)|[a-zA-Z][.)]|"
    r"(?:סעיף|Section|Article|Clause|פסקה)\s+\d+(?:\.\d+)*)(?=\s)\s*"
)
_BULLET_RE = re.compile(r"^\s*[-•●▪*·◦]\s+")
_QUESTION_RE = re.compile(r"^\s*(?:ש|Q|שאלה)\s*[.:]\s*")
_ANSWER_RE = re.compile(r"^\s*(?:ת|A|תשובה)\s*[.:]\s*")
# "עו"ד כהן:", "כב' השופטת:", "העד:", "THE COURT:", "MR. SMITH:" -- in transcripts only.
_SPEAKER_RE = re.compile(r"^\s*(?P<speaker>[^\s:|]{1,25}(?:\s[^\s:|]{1,25}){0,3})\s*:\s+\S")
_EMAIL_SPLIT_RE = re.compile(
    r"^\s*(?:-{2,}\s*(?:Original Message|Forwarded message|הודעה מקורית|הודעה שהועברה)\s*-{2,}|"
    r"(?:From|מאת)\s*:.+|On .{6,120} wrote:|ב[-־]?.{4,80}\s(?:כתב|כתבה|נכתב)\s*:)\s*$",
    re.IGNORECASE,
)
_TERMINAL = tuple(".!?:;\"'”״)]")


@dataclass
class Unit:
    text: str
    page: int | None
    kind: str = "para"  # para | heading | qa
    label: str = ""     # paragraph / clause number, when it has one
    section: str = ""   # the heading in force where the unit starts
    tokens: int = 0


def _is_heading(line: str, previous: str | None) -> bool:
    stripped = line.strip()
    if stripped.startswith("## "):
        return True
    if not (2 <= len(stripped) <= 70) or _EMAIL_HEADER_RE.match(stripped) or _EMAIL_SPLIT_RE.match(stripped):
        return False
    numbered = _NUMBERED_RE.match(stripped)
    if numbered and (len(stripped.split()) > 6 or stripped.endswith(tuple(".;,"))):
        return False  # a short numbered paragraph, not a numbered heading ("3. Background")
    if previous is not None and previous.strip() and not previous.strip().endswith(_TERMINAL):
        return False  # mid-paragraph
    letters = re.sub(r"[^A-Za-zא-ת]", "", stripped)
    if len(letters) < 2:
        return False
    if stripped.endswith(":") and len(stripped) <= 50:
        return True
    return not stripped.endswith(tuple(".,;!?\"'”״")) and (stripped.isupper() or len(stripped.split()) <= 8)


def _starts_unit(line: str, transcript: bool) -> bool:
    return bool(_NUMBERED_RE.match(line) or _BULLET_RE.match(line) or _QUESTION_RE.match(line)
                or _ANSWER_RE.match(line) or _EMAIL_HEADER_RE.match(line) or _EMAIL_SPLIT_RE.match(line)
                or (transcript and _SPEAKER_RE.match(line)))


def build_units(pages: list[PageText], doc_type: str) -> list[Unit]:
    """The document as paragraphs / clauses / speaker turns, each with its page, number and the
    heading it falls under. Visual lines (PDF, OCR, hard-wrapped text) are rejoined into
    paragraphs: a line continues the paragraph unless a blank line, a number/bullet/speaker marker
    or a heading starts a new one, or the previous line ended a sentence well short of the page's
    usual line width."""
    transcript = doc_type == "transcript"
    units: list[Unit] = []
    section = ""
    for page in pages:
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in page.text.splitlines()]
        lengths = sorted(len(line) for line in lines if line)
        full_width = lengths[int(len(lengths) * 0.9)] if lengths else 0
        current: list[str] = []
        previous: str | None = None

        def flush() -> None:
            if current:
                text = " ".join(current).strip()
                marker = _NUMBERED_RE.match(text)
                label = marker.group("label").strip("().") if marker else ""
                units.append(Unit(text, page.number, "para", label, section))
                current.clear()

        for line in lines:
            if not line:
                flush()
                previous = None
                continue
            if _is_heading(line, previous) and not (transcript and _SPEAKER_RE.match(line)):
                flush()
                section = re.sub(r"^#+\s*", "", line).rstrip(":").strip()
                units.append(Unit(section, page.number, "heading", "", section))
                previous = line
                continue
            new = (not page.wrapped or not current or _starts_unit(line, transcript)
                   or (previous is not None and previous.endswith(_TERMINAL) and len(previous) < 0.85 * full_width))
            if new:
                flush()
            current.append(line)
            previous = line
        flush()
    if transcript:
        units = _pair_questions(units)
    return units


def _pair_questions(units: list[Unit]) -> list[Unit]:
    """A question kept with the answer right after it, so testimony is never cut between them."""
    out: list[Unit] = []
    for unit in units:
        if out and out[-1].kind == "para" and _QUESTION_RE.match(out[-1].text) and _ANSWER_RE.match(unit.text):
            out[-1] = Unit(f"{out[-1].text}\n{unit.text}", out[-1].page, "qa", out[-1].label, out[-1].section)
        else:
            out.append(unit)
    return out


_SENTENCE_RE = re.compile(r"(?<=[.;!?])\s+")


def _split_unit(unit: Unit, budget: int) -> list[Unit]:
    """A unit longer than the budget, as sentence-packed pieces (hard-cut as a last resort)."""
    pieces: list[str] = []
    current, tokens = [], 0
    for sentence in _SENTENCE_RE.split(unit.text):
        size = count_tokens(sentence)
        if size > budget:
            if current:
                pieces.append(" ".join(current))
                current, tokens = [], 0
            step = max(budget * 2, 200)  # ~2 characters per token for Hebrew, conservatively
            pieces += [sentence[i : i + step] for i in range(0, len(sentence), step)]
            continue
        if current and tokens + size > budget:
            pieces.append(" ".join(current))
            current, tokens = [], 0
        current.append(sentence)
        tokens += size
    if current:
        pieces.append(" ".join(current))
    return [Unit(p, unit.page, unit.kind, unit.label, unit.section, count_tokens(p)) for p in pieces if p.strip()]


def pack_units(units: list[Unit], budget: int, overlap: int) -> list[list[Unit]]:
    """Whole units packed into groups of at most `budget` tokens. A heading starts a new group
    (once the current one holds something worth keeping on its own), and each group after the
    first opens with the previous group's last units, up to `overlap` tokens -- never with a
    heading, and never a unit that is itself over the overlap."""
    sized: list[Unit] = []
    for unit in units:
        unit.tokens = unit.tokens or count_tokens(unit.text)
        sized += [unit] if unit.tokens <= budget else _split_unit(unit, budget)
    groups: list[list[Unit]] = []
    current: list[Unit] = []
    tokens = 0
    fresh = 0  # tokens of units not carried over from the previous group

    def close() -> None:
        nonlocal current, tokens, fresh
        if fresh:
            groups.append(current)
        carried: list[Unit] = []
        carried_tokens = 0
        for previous in reversed(current if fresh else []):
            if previous.kind == "heading" or carried_tokens + previous.tokens > overlap:
                break
            carried.insert(0, previous)
            carried_tokens += previous.tokens
        current, tokens, fresh = carried, carried_tokens, 0

    for unit in sized:
        if unit.kind == "heading" and fresh >= budget * 0.35:
            close()
            current, tokens = [], 0  # no overlap into a new section
        elif current and tokens + unit.tokens > budget:
            close()
            while current and tokens + unit.tokens > budget:
                tokens -= current.pop(0).tokens
        current.append(unit)
        tokens += unit.tokens
        fresh += unit.tokens
    if fresh:
        groups.append(current)
    return groups


# --- e-mail threads ---------------------------------------------------------------------------


def split_messages(units: list[Unit]) -> list[tuple[str, list[Unit]]]:
    """An e-mail thread as (message header, units) per message: a new message starts at a
    "From:" / "-----Original Message-----" / "On ... wrote:" line. The header is the message's
    From/To/Date/Subject lines, repeated in each of its chunks."""
    messages: list[tuple[list[str], list[Unit]]] = [([], [])]
    in_header = True
    for unit in units:
        first_line = unit.text.splitlines()[0] if unit.text else ""
        if _EMAIL_SPLIT_RE.match(first_line) and (messages[-1][1] or not in_header):
            messages.append(([], []))
            in_header = True
        if in_header and (_EMAIL_HEADER_RE.match(unit.text) or _EMAIL_SPLIT_RE.match(first_line)):
            for line in unit.text.splitlines():
                if _EMAIL_HEADER_RE.match(line):
                    messages[-1][0].append(line.strip())
            continue
        in_header = False
        messages[-1][1].append(unit)
    out = []
    for header_lines, body in messages:
        if body:
            fields = [line for line in header_lines
                      if re.match(r"\s*(From|To|Date|Sent|Subject|מאת|אל|תאריך|נשלח|נושא)\s*:", line, re.IGNORECASE)]
            out.append((" ; ".join(fields)[:400], body))
    return out


# --- chunks -----------------------------------------------------------------------------------

_TYPE_LABELS = {
    "email": "e-mail", "transcript": "transcript / protocol", "judgment": "judgment / decision",
    "affidavit": "affidavit", "pleading": "pleading / motion", "contract": "contract",
    "correspondence": "letter", "evidence": "exhibit / evidence", "notes": "memo / notes", "document": "document",
}


def file_id(rel_path: str) -> str:
    return hashlib.sha1(rel_path.encode("utf-8")).hexdigest()[:12]


def _range(values: list) -> str:
    values = [v for v in values if v not in (None, "")]
    if not values:
        return ""
    first, last = values[0], values[-1]
    return str(first) if first == last else f"{first}-{last}"


def context_header(doc: CaseDocument, group: list[Unit], message: str = "") -> str:
    pages = _range(sorted({u.page for u in group if u.page is not None}))
    labels = _range([u.label for u in group if u.label])
    section = next((u.section for u in group if u.section), "")
    parts = [f"Matter: {doc.matter}", f"File: {doc.rel_path}", f"Type: {_TYPE_LABELS[doc.doc_type]}"]
    if doc.doc_date:
        parts.append(f"Date: {doc.doc_date}")
    if pages:
        parts.append(f"{'Pages' if '-' in pages else 'Page'} {pages}")
    if section and section != doc.title:
        parts.append(f"Section: {section[:120]}")
    if labels:
        parts.append(f"{'Q&A' if doc.doc_type == 'transcript' else 'Para.'} {labels}")
    header = " | ".join(parts)
    if doc.title and doc.title not in header:
        header += f"\nTitle: {doc.title}"
    if message:
        header += f"\nE-mail: {message}"
    return header


def chunk_document(doc: CaseDocument, budget: int, overlap: int, file_sha256: str = "") -> list[CorpusChunk]:
    """The document's chunks, each with its context header (see the module docstring)."""
    units = build_units(doc.pages, doc.doc_type)
    if doc.doc_type == "email":
        parts = split_messages(units) or [("", units)]
    else:
        parts = [("", units)]
    fid = file_id(doc.rel_path)
    drafts: list[tuple[str, list[Unit]]] = []
    for message, body in parts:
        # The header's own size comes out of the budget, so a chunk stays within it whole.
        room = max(budget - count_tokens(context_header(doc, body[:1], message)) - 10, 80)
        drafts += [(message, group) for group in pack_units(body, room, overlap)]
    chunks = []
    for index, (message, group) in enumerate(drafts):
        header = context_header(doc, group, message)
        body = "\n".join(f"## {u.text}" if u.kind == "heading" else u.text for u in group)
        text = f"{header}\n\n{body}"
        pages = [u.page for u in group if u.page is not None]
        metadata = {
            "matter": doc.matter, "file": doc.rel_path, "file_name": doc.path.name, "file_id": fid,
            "doc_type": doc.doc_type, "title": doc.title[:300], "doc_date": doc.doc_date,
            "section": next((u.section for u in group if u.section), "")[:300],
            "labels": _range([u.label for u in group if u.label]), "chunk_index": index, "chunk_count": len(drafts),
            "file_sha256": file_sha256, "chunker_version": CHUNKER_VERSION,
        }
        if pages:
            metadata.update(page_start=min(pages), page_end=max(pages))
        if doc.doc_date:
            metadata["doc_ymd"] = int(doc.doc_date.replace("-", ""))
        if message:
            metadata["email"] = message[:300]
        chunks.append(CorpusChunk(f"{fid}:{index}", text, normalize_for_embedding(text), normalize_for_index(text),
                                  metadata))
    return chunks


# --- the index --------------------------------------------------------------------------------


def _paths() -> tuple[Path, Path]:
    cfg = get_config().legal.case_files
    return Path(cfg.folder), Path(cfg.vectordb_dir)


def record_hash(file_sha256: str) -> str:
    """What decides whether a file is re-chunked: its content and the chunking settings."""
    cfg = get_config().legal.case_files
    key = f"{file_sha256}|{CHUNKER_VERSION}|{cfg.chunk_max_tokens}|{cfg.chunk_overlap_tokens}"
    return hashlib.sha256(key.encode("utf-8")).hexdigest()


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


@dataclass
class FileReport:
    rel_path: str
    action: str  # indexed | unchanged | failed | skipped | dry_run | pruned | missing
    message: str = ""
    doc_type: str = ""
    chunks: int = 0
    notes: list[str] = field(default_factory=list)


def _embed(texts: list[str]):
    import numpy as np

    from docslides.rag.embedding import get_embedder

    cfg = get_config().legal
    model = get_embedder(cfg.retrieval.embedding_model, cfg.retrieval.device)
    return np.asarray(model.encode(texts, batch_size=cfg.case_files.embed_batch_size, normalize_embeddings=True,
                                   convert_to_numpy=True), dtype=np.float32)


def open_store():
    from docslides.legal_data.corpus_index import CorpusCollection, CorpusState

    _, vectordb = _paths()
    vectordb.mkdir(parents=True, exist_ok=True)
    return CorpusCollection(vectordb, get_config().legal.case_files.collection), CorpusState(vectordb / STATE_FILE)


def index_folder(dry_run: bool = False, prune: bool = False, force: bool = False,
                 on_file=None) -> list[FileReport]:
    """Indexes new and changed files in legal.case_files.folder; unchanged ones are skipped. With
    `prune`, files deleted from the folder are removed from the index (else only reported).
    `on_file(report)` is called after each file. State is committed after every file, so an
    interrupted run resumes where it stopped."""
    folder, _ = _paths()
    if not folder.is_dir():
        raise FileNotFoundError(f"{folder} does not exist: create it and put the case files in it "
                                "(one subfolder per matter)")
    files, unsupported = list_case_files(folder)
    reports = [FileReport(p.relative_to(folder).as_posix(), "skipped",
                          f"convert to {CONVERT_FIRST[p.suffix.lower()]} first") for p in unsupported]
    collection, state = (None, None) if dry_run else open_store()
    try:
        seen = set()
        for path in files:
            rel = path.relative_to(folder).as_posix()
            seen.add(rel)
            report = _index_file(path, folder, rel, collection, state, dry_run, force)
            reports.append(report)
            if on_file:
                on_file(report)
        if state is not None:
            for rel in sorted(state.record_ids(CATEGORY) - seen):
                if prune:
                    previous = state.get(CATEGORY, rel)
                    collection.delete(previous.chunk_ids)
                    state.delete(CATEGORY, rel)
                    state.commit()
                    reports.append(FileReport(rel, "pruned", f"{len(previous.chunk_ids)} chunks removed"))
                else:
                    reports.append(FileReport(rel, "missing", "file deleted; still indexed (use --prune)"))
    finally:
        if state is not None:
            state.close()
    _lexical_index.cache_clear()
    return reports


def _index_file(path: Path, folder: Path, rel: str, collection, state, dry_run: bool, force: bool) -> FileReport:
    try:
        sha = _sha256(path)
        rhash = record_hash(sha)
        previous = state.get(CATEGORY, rel) if state is not None else None
        if previous and previous.record_hash == rhash and not force:
            return FileReport(rel, "unchanged", chunks=len(previous.chunk_ids))
        cfg = get_config().legal.case_files
        doc = load_document(path, folder)
        chunks = chunk_document(doc, cfg.chunk_max_tokens, cfg.chunk_overlap_tokens, sha)
        if not chunks:
            return FileReport(rel, "failed", "no readable text", doc.doc_type, notes=doc.notes)
        if dry_run:
            return FileReport(rel, "dry_run", doc.title, doc.doc_type, len(chunks), doc.notes)
        vectors = _embed([c.embed_text for c in chunks])
        if previous:
            stale = set(previous.chunk_ids) - {c.chunk_id for c in chunks}
            collection.delete(sorted(stale))
        collection.upsert(chunks, vectors)
        state.put(CATEGORY, rel, rhash, chunks)
        state.commit()
        return FileReport(rel, "indexed", doc.title, doc.doc_type, len(chunks), doc.notes)
    except Exception as exc:  # noqa: BLE001 -- one bad file mustn't block the rest
        logger.warning("case_files_index_failed", path=rel, error=str(exc))
        return FileReport(rel, "failed", str(exc))


def preview(path: Path) -> list[CorpusChunk]:
    """A file's chunks without indexing anything (scripts/ingest_case_files.py show)."""
    folder, _ = _paths()
    cfg = get_config().legal.case_files
    path, folder = path.resolve(), folder.resolve()
    root = folder if path.is_relative_to(folder) else path.parent
    return chunk_document(load_document(path, root), cfg.chunk_max_tokens, cfg.chunk_overlap_tokens)


# --- retrieval --------------------------------------------------------------------------------


@dataclass
class Excerpt:
    chunk_id: str
    text: str
    metadata: dict
    score: float = 0.0
    via: str = ""  # dense / lexical / both / neighbor

    def citation(self) -> str:
        m = self.metadata
        where = m.get("file", "")
        start, end = m.get("page_start"), m.get("page_end")
        if start:
            where += f", p. {start}" if start == end else f", pp. {start}-{end}"
        if m.get("labels"):
            where += f", {'Q&A' if m.get('doc_type') == 'transcript' else 'para.'} {m['labels']}"
        return where


@lru_cache(maxsize=8)
def _lexical_index(matter: str | None):
    from docslides.legal.corpus_lexical import CorpusLexicalIndex
    from docslides.legal_data.corpus_index import CorpusState

    _, vectordb = _paths()
    state = CorpusState(vectordb / STATE_FILE)
    try:
        rows = [(cid, text) for cid, rel, text in state.lexical(CATEGORY) if matter is None or matter_of(rel) == matter]
    finally:
        state.close()
    return CorpusLexicalIndex.build([r[0] for r in rows], [r[1] for r in rows]) if rows else None


def matters() -> list[str]:
    from docslides.legal_data.corpus_index import CorpusState

    _, vectordb = _paths()
    if not (vectordb / STATE_FILE).exists():
        return []
    state = CorpusState(vectordb / STATE_FILE)
    try:
        return sorted({matter_of(rel) for rel in state.record_ids(CATEGORY)})
    finally:
        state.close()


def _rerank(query: str, excerpts: list[Excerpt]) -> list[Excerpt]:
    retrieval_cfg = get_config().legal.retrieval
    if not retrieval_cfg.reranker_model or not excerpts:
        return excerpts
    from docslides.legal.retrieval import _reranker, reranker_device

    model = _reranker(retrieval_cfg.reranker_model, reranker_device())
    if model is None:
        return excerpts
    scores = model.predict([(query, e.text) for e in excerpts])
    for excerpt, score in zip(excerpts, scores):
        excerpt.score = float(score)
    return sorted(excerpts, key=lambda e: -e.score)


def search(queries: list[str], matter: str | None = None, doc_type: str | None = None) -> list[Excerpt]:
    """The best excerpts for `queries` (the question, and for a follow-up also the question with
    the previous one): dense and BM25 rankings fused (reciprocal rank), reranked against the first
    query, cut to top_k, then the chunks before and after the best hits added -- all within
    max_evidence_tokens. Returned in document order (file, then position)."""
    cfg = get_config().legal.case_files
    collection, state = open_store()
    state.close()
    total = collection.count()
    if total == 0:
        return []
    filters = [{"matter": matter}] if matter else []
    if doc_type:
        filters.append({"doc_type": doc_type})
    where = filters[0] if len(filters) == 1 else ({"$and": filters} if filters else None)

    fused: dict[str, float] = {}
    via: dict[str, set[str]] = {}
    found: dict[str, tuple[str, dict]] = {}
    vectors = _embed([normalize_for_embedding(q) for q in queries])
    for query, vector in zip(queries, vectors):
        result = collection.query(vector.tolist(), min(cfg.fetch_k, total), where)
        for rank, (cid, document, meta) in enumerate(zip(result["ids"][0], result["documents"][0],
                                                         result["metadatas"][0])):
            fused[cid] = fused.get(cid, 0.0) + 1.0 / (60 + rank)
            via.setdefault(cid, set()).add("dense")
            found[cid] = (document, meta)
        lexical = _lexical_index(matter)
        for rank, (cid, _score) in enumerate(lexical.search(query, cfg.fetch_k) if lexical else []):
            fused[cid] = fused.get(cid, 0.0) + 1.0 / (60 + rank)
            via.setdefault(cid, set()).add("lexical")
    missing = [cid for cid in fused if cid not in found]
    if missing:
        got = collection.collection.get(ids=missing, include=["documents", "metadatas"])
        found.update({cid: (doc, meta) for cid, doc, meta in zip(got["ids"], got["documents"], got["metadatas"])})
    ranked = [Excerpt(cid, found[cid][0], found[cid][1], fused[cid], "both" if len(via[cid]) > 1 else next(iter(via[cid])))
              for cid in sorted(fused, key=lambda c: -fused[c]) if cid in found]
    if doc_type:
        ranked = [e for e in ranked if e.metadata.get("doc_type") == doc_type]
    ranked = _rerank(queries[0], ranked[: cfg.rerank_candidates])[: cfg.top_k]

    budget = cfg.max_evidence_tokens
    chosen: list[Excerpt] = []
    for excerpt in ranked:
        size = count_tokens(excerpt.text)
        if size <= budget or not chosen:
            chosen.append(excerpt)
            budget -= size
    have = {e.chunk_id for e in chosen}
    neighbor_ids = []
    for excerpt in chosen[: cfg.neighbor_hits]:
        m = excerpt.metadata
        for index in (m["chunk_index"] - 1, m["chunk_index"] + 1):
            cid = f"{m['file_id']}:{index}"
            if 0 <= index < m["chunk_count"] and cid not in have:
                neighbor_ids.append(cid)
                have.add(cid)
    if neighbor_ids:
        got = collection.collection.get(ids=neighbor_ids, include=["documents", "metadatas"])
        for cid, document, meta in zip(got["ids"], got["documents"], got["metadatas"]):
            size = count_tokens(document)
            if size <= budget:
                chosen.append(Excerpt(cid, document, meta, 0.0, "neighbor"))
                budget -= size
    return sorted(chosen, key=lambda e: (e.metadata.get("file", ""), e.metadata.get("chunk_index", 0)))


# --- answering --------------------------------------------------------------------------------

SYSTEM_PROMPT = """You are a legal assistant helping an attorney work with the documents in their own \
case files. You answer ONLY from the numbered excerpts in <excerpts>; each excerpt starts with a header \
naming its matter, file, document type, date, pages and paragraph numbers.

Rules:
- End every factual statement with the number(s) of the excerpt(s) it comes from, like [2] or [1][4].
- Copy names, dates, amounts, case numbers, and paragraph/clause numbers exactly as written.
- Attribute positions. A statement in a pleading, affidavit, letter or e-mail is that party's claim, \
not an established fact: say who asserts it ("the plaintiff alleges ... [3]"). Distinguish what a court \
decided (judgments, decisions) from what the parties argue.
- When documents disagree, give each version with its source.
- When asked for a chronology, order events by date and cite each one.
- If the excerpts do not answer the question, say so plainly and say what is missing (which document \
or fact you would need). Never fill gaps about this case from outside knowledge.
- General legal explanations only when asked for, marked as general and not from the file.
- Answer in the language of the question; keep quotations in their original language."""


def format_excerpts(excerpts: list[Excerpt]) -> str:
    return "\n\n".join(f"[{n}] {e.text}" for n, e in enumerate(excerpts, 1))


def cited(answer: str, excerpts: list[Excerpt]) -> list[tuple[int, Excerpt]]:
    numbers = sorted({int(n) for n in re.findall(r"\[(\d{1,3})\]", answer)})
    return [(n, excerpts[n - 1]) for n in numbers if 1 <= n <= len(excerpts)]


@dataclass
class Turn:
    question: str
    answer: str


async def answer(question: str, history: list[Turn] | None = None, matter: str | None = None,
                 doc_type: str | None = None) -> tuple[str, list[Excerpt]]:
    """A grounded answer to `question` from the case files, and the excerpts it was given."""
    from docslides.llm.client import ChatMessage, LLMCallSite, SamplingParams, get_legal_orchestrator_client

    cfg = get_config().legal.case_files
    history = (history or [])[-cfg.history_turns:] if cfg.history_turns else []
    queries = [question] + ([f"{history[-1].question}\n{question}"] if history else [])
    excerpts = await asyncio.to_thread(search, queries, matter, doc_type)
    if not excerpts:
        return ("Nothing in the indexed case files matches this question"
                + (f" in matter '{matter}'" if matter else "") + ". Check that the files are in "
                f"{cfg.folder} and indexed (python scripts/ingest_case_files.py).", [])
    messages = [ChatMessage("system", SYSTEM_PROMPT)]
    for turn in history:
        messages += [ChatMessage("user", turn.question), ChatMessage("assistant", turn.answer)]
    messages.append(ChatMessage("user", f"<excerpts>\n{format_excerpts(excerpts)}\n</excerpts>\n\n"
                                        f"Question: {question}"))
    reply = await get_legal_orchestrator_client().complete_text(
        messages, LLMCallSite("legal_case_files_chat"),
        sampling=SamplingParams(temperature=0.2, top_p=0.9, max_tokens=cfg.answer_max_tokens, seed=0),
        enable_thinking=False,
    )
    return reply, excerpts
