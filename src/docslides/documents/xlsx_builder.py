"""Build an .xlsx workbook from a `SpreadsheetSpec` with openpyxl: a bold,
frozen header row, numbers stored as numbers, '='-prefixed strings as Excel
formulas, columns sized to their content, and right-to-left sheets for RTL
languages."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from docslides.config import get_config
from docslides.llm.schemas import CellValue, SpreadsheetSpec
from docslides.logging_setup import get_logger

logger = get_logger(__name__)

_SHEET_NAME_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
_NUMBER = re.compile(r"[-+]?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?")
_HEADER_FILL = PatternFill("solid", fgColor="F97316")
_MAX_COLUMN_WIDTH = 60


def _sheet_name(name: str, taken: set[str]) -> str:
    base = _SHEET_NAME_FORBIDDEN.sub(" ", name).strip().strip("'")[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in taken:
        suffix = f" ({n})"
        candidate, n = base[: 31 - len(suffix)] + suffix, n + 1
    taken.add(candidate.lower())
    return candidate


def _cell_value(value: CellValue) -> CellValue | int:
    """Numbers the model wrote as text ("1,250", "3.5") are stored as numbers, so they sum."""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if not isinstance(value, str):
        return value
    text = ILLEGAL_CHARACTERS_RE.sub("", value).strip()
    if _NUMBER.fullmatch(text):
        number = float(text.replace(",", ""))
        return int(number) if number.is_integer() and "." not in text else number
    return text


def build_xlsx(spec: SpreadsheetSpec, output_path: str | Path, lang: str = "en") -> Path:
    rtl = get_config().languages.is_rtl(lang)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = spec.title
    taken: set[str] = set()

    for sheet_spec in spec.sheets:
        sheet = workbook.create_sheet(_sheet_name(sheet_spec.name, taken))
        sheet.sheet_view.rightToLeft = rtl
        widths = [len(str(c)) for c in sheet_spec.columns]

        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", str(c)) for c in sheet_spec.columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in sheet_spec.rows:
            values = [_cell_value(v) for v in row[: max(len(sheet_spec.columns), 1)]]
            sheet.append(values)
            for i, value in enumerate(values[: len(widths)]):
                if not (isinstance(value, str) and value.startswith("=")):  # a formula's text isn't shown
                    widths[i] = max(widths[i], len(str(value if value is not None else "")))

        for i, width in enumerate(widths, start=1):
            column = sheet.column_dimensions[get_column_letter(i)]
            column.width = min(max(width + 2, 8), _MAX_COLUMN_WIDTH)
        sheet.freeze_panes = "A2"
        if sheet_spec.columns and sheet_spec.rows:
            sheet.auto_filter.ref = sheet.dimensions

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(output_path))
    logger.info("xlsx_built", output=str(output_path), sheets=len(spec.sheets))
    return output_path
