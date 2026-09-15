"""Explicit RTL support for Arabic/Hebrew slides.

python-pptx's high-level API has no notion of right-to-left paragraph
direction, so we manipulate the underlying OOXML directly via lxml:
  * set `rtl="1"` on every `<a:pPr>` (paragraph properties) element in an
    RTL text frame,
  * right-align RTL text boxes,
  * mirror the horizontal position of asymmetric layout elements (title
    bars, footer icons, decorative accents) so a deck built from a
    Western-oriented template still reads correctly in RTL.

Supports mixed-direction decks: RTL treatment is applied per-slide, not
globally, so an English slide and an Arabic slide can coexist in one deck.
"""

from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.shapes.base import BaseShape
from pptx.slide import Slide
from pptx.text.text import TextFrame

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _qn(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


def set_textframe_rtl(text_frame: TextFrame) -> None:
    """Set rtl="1" on every paragraph's <a:pPr> and right-align the text."""
    for paragraph in text_frame.paragraphs:
        p_elem = paragraph._p  # CT_TextParagraph, the underlying lxml element
        pPr = p_elem.get_or_add_pPr()
        pPr.set("rtl", "1")
        paragraph.alignment = PP_ALIGN.RIGHT


def apply_rtl_to_shape(shape: BaseShape) -> None:
    if shape.has_text_frame:
        set_textframe_rtl(shape.text_frame)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                set_textframe_rtl(cell.text_frame)


def mirror_shape_horizontal(shape: BaseShape, slide_width_emu: int) -> None:
    """Mirror an asymmetric layout element's horizontal position (e.g. a title
    bar or footer icon anchored to one side) so it lands on the correct side
    for RTL reading order."""
    if shape.left is None or shape.width is None:
        return
    shape.left = slide_width_emu - shape.left - shape.width


def apply_rtl_to_slide(
    slide: Slide,
    slide_width_emu: int,
    mirror_shape_names: set[str] | None = None,
) -> None:
    """Apply full RTL treatment to one slide: paragraph direction + alignment
    on every text-bearing shape, and horizontal mirroring for named
    asymmetric layout elements (pass their shape `.name` from the template,
    e.g. {"Title Bar", "Footer Icon"})."""
    mirror_shape_names = mirror_shape_names or set()
    for shape in slide.shapes:
        apply_rtl_to_shape(shape)
        if shape.name in mirror_shape_names:
            mirror_shape_horizontal(shape, slide_width_emu)
